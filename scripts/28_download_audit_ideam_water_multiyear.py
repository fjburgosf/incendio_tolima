"""Descarga y audita una tabla pública de promedios multianuales de calidad del agua."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from urllib.request import urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
URL = "https://bart.ideam.gov.co/indiecosistemas/ind/agua/datos/D_Promedio%20multianual%20variables%20calidad%20agua.xlsx"
RAW_DIR = ROOT / "data" / "raw" / "ideam_water"
OUTPUT = RAW_DIR / "ideam_water_multiyear_variables.xlsx"
MANIFEST = RAW_DIR / "ideam_water_multiyear_variables_manifest.json"
AUDIT = ROOT / "results" / "exploratory" / "audit_ideam_water_multiyear_variables.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def nonempty_rows(sheet, limit: int = 12) -> list[list[object]]:
    rows: list[list[object]] = []
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        if any(value is not None for value in values):
            rows.append(values)
        if len(rows) >= limit:
            break
    return rows


def main() -> None:
    if OUTPUT.exists() or MANIFEST.exists() or AUDIT.exists():
        raise FileExistsError("La auditoría de calidad del agua ya existe. No se sobreescriben insumos ni resultados.")
    with urlopen(URL, timeout=120) as response:
        payload = response.read()
        content_type = response.headers.get_content_type()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_bytes(payload)
    manifest = {
        "url": URL,
        "file": OUTPUT.name,
        "bytes": len(payload),
        "sha256": sha256(OUTPUT),
        "content_type": content_type,
        "purpose": "Auditoría de estructura. No es una serie de calidad del agua ni un análisis del evento.",
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    workbook = load_workbook(OUTPUT, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        samples = nonempty_rows(sheet)
        sheets.append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "first_nonempty_rows": samples,
        })
    workbook.close()
    audit = {
        "source_file": str(OUTPUT.relative_to(ROOT)).replace("\\", "/"),
        "sheet_count": len(sheets),
        "sheets": sheets,
        "interpretation_limit": (
            "El archivo describe promedios multianuales. No es una serie temporal diaria o mensual, "
            "no demuestra condición posterior al incendio y no permite atribución causal."
        ),
    }
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({
        "bytes": manifest["bytes"],
        "sha256": manifest["sha256"],
        "sheet_count": audit["sheet_count"],
        "sheets": [{"name": row["name"], "max_row": row["max_row"], "max_column": row["max_column"]} for row in sheets],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
