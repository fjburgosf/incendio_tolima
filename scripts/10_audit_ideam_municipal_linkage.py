"""Audita el enlace nominal entre municipios de eventos IDEAM y Cortolima.

No modifica ni combina archivos. El enlace es una comprobación de nombres para
una posible agregación municipal y no identifica un evento individual.
"""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCES = (
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2002_2010.xlsx",
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2010_2020.xlsx",
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2021_2023_2024_0.xlsx",
)
BOUNDARIES = ROOT / "data" / "raw" / "boundaries" / "cortolima_limite_municipal_25k.geojson"
CROSSWALK = ROOT / "configs" / "ideam_cortolima_municipality_crosswalk.json"
OUTPUT = ROOT / "results" / "exploratory" / "audit_ideam_municipal_linkage.json"


def normalized(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]+", "", text.upper())


def find_header(worksheet) -> tuple[int, list[object]]:
    for number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        values = [normalized(value) for value in row]
        if "DEPARTAMENTO" in values and "MUNICIPIO" in values and any(value.startswith("DIA") for value in values):
            return number, list(row)
    raise ValueError(f"No se encontró encabezado de datos en {worksheet.title}")


def main() -> None:
    boundary_data = json.loads(BOUNDARIES.read_text(encoding="utf-8"))
    crosswalk_data = json.loads(CROSSWALK.read_text(encoding="utf-8"))
    boundary_names = {
        normalized(feature["properties"].get("mun_nombre")): feature["properties"].get("mun_nombre")
        for feature in boundary_data["features"]
    }
    crosswalk = crosswalk_data["mappings"]
    sheets = []
    for source in SOURCES:
        workbook = load_workbook(source, read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            try:
                header_row, header = find_header(worksheet)
            except ValueError:
                continue
            names = [normalized(value) for value in header]
            department_column = names.index("DEPARTAMENTO")
            municipality_column = names.index("MUNICIPIO")
            municipalities: dict[str, dict[str, object]] = {}
            for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                department = normalized(row[department_column] if department_column < len(row) else None)
                municipality = row[municipality_column] if municipality_column < len(row) else None
                municipality_key = normalized(municipality)
                if department == "TOLIMA" and municipality_key:
                    entry = municipalities.setdefault(
                        municipality_key,
                        {"original": str(municipality).strip(), "record_count": 0},
                    )
                    entry["record_count"] += 1
            direct_matches = sorted(
                {entry["original"] for key, entry in municipalities.items() if key in boundary_names},
                key=normalized,
            )
            crosswalk_matches = sorted(
                {
                    entry["original"]
                    for key, entry in municipalities.items()
                    if key not in boundary_names
                    and key in crosswalk
                    and normalized(crosswalk[key]["cortolima_municipality"]) in boundary_names
                },
                key=normalized,
            )
            unresolved = sorted(
                {
                    entry["original"]
                    for key, entry in municipalities.items()
                    if key not in boundary_names and key not in crosswalk
                },
                key=normalized,
            )
            sheets.append(
                {
                    "source_book": source.name,
                    "sheet": worksheet.title,
                    "unique_tolima_municipality_name_count": len(municipalities),
                    "unique_direct_normalized_match_count": len(direct_matches),
                    "unique_crosswalk_match_count": len(crosswalk_matches),
                    "crosswalk_ideam_municipalities": crosswalk_matches,
                    "unresolved_ideam_municipalities": unresolved,
                    "unresolved_record_count": sum(
                        entry["record_count"]
                        for key, entry in municipalities.items()
                        if key not in boundary_names and key not in crosswalk
                    ),
                }
            )
    audit = {
        "boundary_source": str(BOUNDARIES.relative_to(ROOT)).replace("\\", "/"),
        "crosswalk_source": str(CROSSWALK.relative_to(ROOT)).replace("\\", "/"),
        "boundary_unique_municipality_name_count": len(boundary_names),
        "name_normalization": "Mayúsculas sin tildes ni caracteres no alfanuméricos",
        "sheets": sheets,
        "interpretation_limit": (
            "Una coincidencia nominal valida una posible llave municipal. No confirma que la fila "
            "describa un incendio individual, un perímetro o una ausencia de evento en otras fechas."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
