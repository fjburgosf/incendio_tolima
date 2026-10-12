"""Audit downloaded public institutional source content without modifying raw files."""

from __future__ import annotations

import json
import re
from html.parser import HTMLParser
from pathlib import Path

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    import openpyxl
except Exception:
    openpyxl = None

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "institutional_fire_records_2026"
MANIFEST = ROOT / "results" / "exploratory" / "audit_public_institutional_sources_download_2026-09-08.json"
OUT = ROOT / "results" / "exploratory" / "audit_public_institutional_content_2026-09-08.json"


class TextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag.lower() in {"script", "style", "noscript"}:
            self.skip += 1

    def handle_endtag(self, tag):
        if tag.lower() in {"script", "style", "noscript"} and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip and data.strip():
            self.parts.append(" ".join(data.split()))


def html_audit(path: Path) -> dict:
    parser = TextParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    text = " ".join(parser.parts)
    title = parser.parts[0] if parser.parts else ""
    return {
        "file": str(path.relative_to(ROOT)),
        "title_prefix": title[:300],
        "published_matches": re.findall(r"(?:Publicado|publicado|Fecha)[^<]{0,80}", text)[:5],
        "hectare_mentions": re.findall(r"[0-9][0-9.,]*\s*(?:hectáreas|has)", text, flags=re.IGNORECASE)[:20],
        "fire_count_mentions": re.findall(r"[0-9][0-9.,]*\s+incendios", text, flags=re.IGNORECASE)[:20],
        "text_characters": len(text),
    }


def pdf_audit(path: Path) -> dict:
    item = {"file": str(path.relative_to(ROOT)), "bytes": path.stat().st_size}
    if PdfReader is None:
        item["status"] = "pypdf_unavailable"
        return item
    try:
        reader = PdfReader(str(path))
        first = reader.pages[0].extract_text() or "" if reader.pages else ""
        item.update(
            {
                "status": "pass",
                "pages": len(reader.pages),
                "first_page_prefix": " ".join(first.split())[:500],
                "tolima_mentions": len(re.findall(r"Tolima", first, flags=re.IGNORECASE)),
            }
        )
    except Exception as exc:
        item.update({"status": "error", "error": repr(exc)})
    return item


def xlsx_audit(path: Path) -> dict:
    item = {"file": str(path.relative_to(ROOT)), "bytes": path.stat().st_size}
    if openpyxl is None:
        item["status"] = "openpyxl_unavailable"
        return item
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet in book.worksheets:
            sheets.append({"title": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column})
        item.update({"status": "pass", "sheets": sheets})
    except Exception as exc:
        item.update({"status": "error", "error": repr(exc)})
    return item


def main() -> None:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    result = {
        "audit_date": "2026-09-08",
        "inference_class": "descriptive source content audit",
        "html": [html_audit(ROOT / x["file"]) for x in manifest["pages"] if x.get("status") == 200],
        "baicv_pdf": [pdf_audit(ROOT / x["file"]) for x in manifest["baicv"] if x.get("status") == 200],
        "ideam_historical_xlsx": [xlsx_audit(ROOT / x["file"]) for x in manifest["ideam_historical"] if x.get("status") == 200],
        "interpretation_limit": "The audit verifies downloaded content and file structure. It does not establish independent perimeter accuracy or equivalence with GWIS.",
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"html": len(result["html"]), "baicv_pdf": len(result["baicv_pdf"]), "xlsx": len(result["ideam_historical_xlsx"]), "output": str(OUT)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
