"""Descarga e inspecciona las estadísticas municipales de transición de MapBiomas."""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path
from urllib.request import urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
URL = "https://storage.googleapis.com/mapbiomas-public/initiatives/colombia/collection_3/statistics/statistics_for_website_mb_colombia_transision_col3.zip"
RAW_DIR = ROOT / "data" / "raw" / "mapbiomas"
OUTPUT = RAW_DIR / "mapbiomas_colombia_collection3_transition_statistics.zip"
MANIFEST = RAW_DIR / "mapbiomas_colombia_collection3_transition_statistics_manifest.json"
AUDIT = ROOT / "results" / "exploratory" / "audit_mapbiomas_colombia_transition_statistics.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def main() -> None:
    if OUTPUT.exists() or MANIFEST.exists() or AUDIT.exists():
        raise FileExistsError("La auditoría de transiciones MapBiomas ya existe. No se sobreescriben insumos ni resultados.")
    with urlopen(URL, timeout=240) as response:
        payload = response.read()
        content_type = response.headers.get_content_type()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_bytes(payload)
    manifest = {"url": URL, "file": OUTPUT.name, "bytes": len(payload), "sha256": sha256(OUTPUT), "content_type": content_type}
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with zipfile.ZipFile(OUTPUT) as archive:
        members = archive.namelist()
        books = [name for name in members if name.lower().endswith(".xlsx")]
        if not books:
            raise ValueError("El ZIP no contiene libro XLSX")
        summaries = []
        for member in books:
            with archive.open(member) as binary:
                workbook = load_workbook(io.BytesIO(binary.read()), read_only=True, data_only=True)
                for sheet in workbook.worksheets:
                    rows = sheet.iter_rows(values_only=True)
                    try:
                        header = [str(value).strip() if value is not None else "" for value in next(rows)]
                    except StopIteration:
                        continue
                    if "departamento" not in header:
                        continue
                    department_index = header.index("departamento")
                    municipality_index = header.index("municipio") if "municipio" in header else None
                    total = 0
                    tolima = 0
                    municipalities: set[str] = set()
                    for row in rows:
                        if not any(value is not None for value in row):
                            continue
                        total += 1
                        if str(row[department_index]).strip().upper() == "TOLIMA":
                            tolima += 1
                            if municipality_index is not None:
                                municipalities.add(str(row[municipality_index]).strip())
                    summaries.append({"member": member, "sheet": sheet.title, "fields": header, "total_records": total, "tolima_record_count": tolima, "tolima_unique_municipality_count": len(municipalities)})
                workbook.close()
    audit = {"source_file": str(OUTPUT.relative_to(ROOT)).replace("\\", "/"), "zip_members": members, "sheets": summaries, "interpretation_limit": "Las transiciones son áreas clasificadas entre periodos. No identifican predios, fechas diarias, incendios o causas de cambio."}
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"bytes": manifest["bytes"], "sha256": manifest["sha256"], "sheets": summaries}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
