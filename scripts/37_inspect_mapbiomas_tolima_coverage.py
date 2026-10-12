"""Inspecciona la tabla municipal de cobertura MapBiomas para Tolima."""

from __future__ import annotations

import io
import json
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "mapbiomas" / "mapbiomas_colombia_collection3_coverage_statistics.zip"
OUTPUT = ROOT / "results" / "exploratory" / "inspect_mapbiomas_tolima_coverage.json"
MEMBER = "Statistics-for-Website-MB-Cobertura-col3.xlsx"
SHEET = "COBERTURA_MUNICIPIO"


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La inspección MapBiomas de Tolima ya existe. No se sobreescriben resultados.")
    with zipfile.ZipFile(SOURCE) as archive, archive.open(MEMBER) as binary:
        workbook = load_workbook(io.BytesIO(binary.read()), read_only=True, data_only=True)
        sheet = workbook[SHEET]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(rows)]
        department_index = header.index("departamento")
        municipality_index = header.index("municipio")
        level0_index = header.index("class_level_0")
        level1_index = header.index("class_level_1")
        level2_index = header.index("class_level_2")
        year_columns = {year: header.index(year) for year in ("1985", "2000", "2010", "2020", "2024")}
        records = 0
        municipalities: set[str] = set()
        classes0: Counter[str] = Counter()
        classes1: Counter[str] = Counter()
        classes2: Counter[str] = Counter()
        nonmissing_by_year: Counter[str] = Counter()
        for row in rows:
            if str(row[department_index]).strip().upper() != "TOLIMA":
                continue
            records += 1
            municipalities.add(str(row[municipality_index]).strip())
            classes0[str(row[level0_index]).strip()] += 1
            classes1[str(row[level1_index]).strip()] += 1
            classes2[str(row[level2_index]).strip()] += 1
            for year, index in year_columns.items():
                if row[index] is not None:
                    nonmissing_by_year[year] += 1
        workbook.close()
    audit = {
        "source_file": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
        "member": MEMBER,
        "sheet": SHEET,
        "tolima_record_count": records,
        "tolima_unique_municipality_count": len(municipalities),
        "class_level_0_values": sorted(classes0),
        "class_level_1_values": sorted(classes1),
        "class_level_2_values": sorted(classes2),
        "nonmissing_records_by_control_year": dict(sorted(nonmissing_by_year.items())),
        "interpretation_limit": "Las áreas son estadísticas de clasificación de cobertura. No son una medición de uso efectivo, predios ni actividad de quema.",
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
