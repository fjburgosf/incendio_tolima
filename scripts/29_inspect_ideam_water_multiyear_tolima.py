"""Inspecciona la cobertura de Tolima en el archivo multianual de calidad del agua."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "ideam_water" / "ideam_water_multiyear_variables.xlsx"
OUTPUT = ROOT / "results" / "exploratory" / "inspect_ideam_water_multiyear_tolima.json"


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La inspección de Tolima ya existe. No se sobreescriben resultados.")
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["Promedio multianual"]
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    next(rows)
    header = next(rows)
    names = [str(value).strip() if value is not None else "" for value in header]
    records = [dict(zip(names, row)) for row in rows if any(value is not None for value in row)]
    workbook.close()
    tolima = [row for row in records if str(row.get("DEPARTAMENTO", "")).strip().upper() == "TOLIMA"]
    fields = [
        "NOMBRE DEL PUNTO DE MONITOREO", "MUNICIPIO", "CORRIENTE", "LONGITUD", "LATITUD",
        "Periodo años monitoreo", "Numero de años de monitoreo", "Total número de monitoreos por punto",
    ]
    audit = {
        "source_file": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
        "total_records": len(records),
        "tolima_record_count": len(tolima),
        "tolima_municipality_counts": dict(sorted(Counter(str(row.get("MUNICIPIO")) for row in tolima).items())),
        "tolima_records": [{field: row.get(field) for field in fields} for row in tolima],
        "interpretation_limit": (
            "Estos son promedios multianuales por punto y escenario hidrológico. No son mediciones posteriores "
            "al episodio, no representan toda la cuenca y no identifican un efecto del incendio."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({
        "total_records": audit["total_records"],
        "tolima_record_count": audit["tolima_record_count"],
        "tolima_municipality_counts": audit["tolima_municipality_counts"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
