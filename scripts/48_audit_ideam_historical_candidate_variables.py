"""Audita presencia y completitud de campos candidatos en libros históricos IDEAM."""

from __future__ import annotations

import json
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCES = (
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2002_2010.xlsx",
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2010_2020.xlsx",
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2021_2023_2024_0.xlsx",
)
OUTPUT = ROOT / "results" / "exploratory" / "audit_ideam_historical_candidate_variables.json"
KEYWORDS = ("CAUSA", "COBERTURA", "AREA", "TIPO", "CLASE", "ORIGEN", "ACTIVIDAD", "ENTIDAD REPORTA")


def normalized(value: object) -> str:
    value = "" if value is None else str(value)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return " ".join(value.upper().split())


def find_header(sheet) -> tuple[int, list[str]] | None:
    for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        header = [normalized(value) for value in row]
        if "DEPARTAMENTO" in header and "MUNICIPIO" in header and any(value.startswith("DIA") for value in header):
            return number, header
    return None


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La auditoría de variables IDEAM ya existe.")
    summaries = []
    for source in SOURCES:
        workbook = load_workbook(source, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            found = find_header(sheet)
            if found is None:
                continue
            header_row, header = found
            department_index = header.index("DEPARTAMENTO")
            selected = [(index, name) for index, name in enumerate(header) if any(keyword in name for keyword in KEYWORDS)]
            total_tolima = 0
            values = {name: Counter() for _, name in selected}
            missing = Counter()
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if department_index >= len(row) or normalized(row[department_index]) != "TOLIMA":
                    continue
                total_tolima += 1
                for index, name in selected:
                    value = row[index] if index < len(row) else None
                    if value is None or str(value).strip() == "":
                        missing[name] += 1
                    else:
                        values[name][str(value).strip()] += 1
            summaries.append({
                "source_book": source.name,
                "sheet": sheet.title,
                "tolima_record_count": total_tolima,
                "candidate_fields": [{"header": name, "missing_count": missing[name], "unique_nonmissing_value_count": len(values[name]), "top_values": values[name].most_common(12)} for _, name in selected],
            })
        workbook.close()
    result = {"sheets": summaries, "interpretation_limit": "La presencia de un campo no prueba que su definición, codificación, completitud o comparabilidad sea suficiente para análisis. Los registros siguen siendo eventos reportados."}
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
