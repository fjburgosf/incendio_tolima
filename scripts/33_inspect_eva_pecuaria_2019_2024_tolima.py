"""Audita la cobertura de Tolima en cada hoja de la base pecuaria EVA."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "eva" / "eva_pecuaria_2019_2024.xlsx"
OUTPUT = ROOT / "results" / "exploratory" / "inspect_eva_pecuaria_2019_2024_tolima.json"


def text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def locate_header(sheet) -> tuple[int, list[str]]:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [text(value) for value in row]
        if "Departamento" in values and "Municipio" in values and "Año" in values:
            return row_number, values
        if row_number > 20:
            break
    raise ValueError(f"No se encontró encabezado esperado en {sheet.title}")


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La inspección EVA pecuaria de Tolima ya existe. No se sobreescriben resultados.")
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    summaries = []
    for sheet in workbook.worksheets:
        header_row, header = locate_header(sheet)
        dept_index = header.index("Departamento")
        municipality_index = header.index("Municipio")
        year_index = header.index("Año")
        rows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
        total = 0
        tolima = 0
        years: Counter[str] = Counter()
        municipalities: Counter[str] = Counter()
        for row in rows:
            if not any(value is not None for value in row):
                continue
            total += 1
            if text(row[dept_index]).upper() == "TOLIMA":
                tolima += 1
                years[text(row[year_index])] += 1
                municipalities[text(row[municipality_index])] += 1
        summaries.append({
            "sheet": sheet.title,
            "header_row": header_row,
            "columns": header,
            "total_data_records": total,
            "tolima_record_count": tolima,
            "tolima_records_by_year": dict(sorted(years.items())),
            "tolima_unique_municipality_count": len(municipalities),
        })
    workbook.close()
    audit = {
        "source_file": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
        "sheets": summaries,
        "interpretation_limit": (
            "Los inventarios están agregados por municipio y año. No miden prácticas de quema, ubicación de predios "
            "ni exposición a un evento de incendio. La continuidad real por especie debe evaluarse por hoja."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps([{
        "sheet": item["sheet"],
        "total_data_records": item["total_data_records"],
        "tolima_record_count": item["tolima_record_count"],
        "tolima_records_by_year": item["tolima_records_by_year"],
        "tolima_unique_municipality_count": item["tolima_unique_municipality_count"],
    } for item in summaries], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
