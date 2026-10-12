"""Audita la estructura de tres libros históricos IDEAM de incendios 2002 a 2024.

No modifica el libro crudo ni consolida eventos. Los conteos sirven para evaluar
la factibilidad de un desenlace histórico, no para producir resultados del estudio.
"""

from __future__ import annotations

import hashlib
import json
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCES = (
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2002_2010.xlsx",
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2010_2020.xlsx",
    ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2021_2023_2024_0.xlsx",
)
OUTPUT = ROOT / "results" / "exploratory" / "audit_ideam_historical_events_2002_2024.json"


def normalized(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return " ".join(text.upper().split())


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def find_header(worksheet) -> tuple[int, list[object]]:
    for number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        values = [normalized(value) for value in row]
        if "DEPARTAMENTO" in values and "MUNICIPIO" in values and any(value.startswith("DIA") for value in values):
            return number, list(row)
    raise ValueError(f"No se encontró encabezado de datos en {worksheet.title}")


def as_iso_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def main() -> None:
    sheets = []
    total_data_rows = 0
    total_tolima_rows = 0
    books = []
    for source in SOURCES:
        workbook = load_workbook(source, read_only=True, data_only=True)
        book_sheets = []
        for worksheet in workbook.worksheets:
            try:
                header_row, header = find_header(worksheet)
            except ValueError:
                sheet_audit = {
                    "source_book": source.name,
                    "sheet": worksheet.title,
                    "worksheet_max_row": worksheet.max_row,
                    "worksheet_max_column": worksheet.max_column,
                    "status": "omitted_without_event_header",
                }
                sheets.append(sheet_audit)
                book_sheets.append(sheet_audit)
                continue
            names = [normalized(value) for value in header]
            department_column = names.index("DEPARTAMENTO")
            municipality_column = names.index("MUNICIPIO")
            day_column = next(index for index, value in enumerate(names) if value.startswith("DIA"))
            area_columns = [value for value in names if value.startswith("AREA TOTAL")]
            area_column = next((index for index, value in enumerate(names) if value.startswith("AREA TOTAL")), None)
            latitude_column = next((index for index, value in enumerate(names) if value == "LATITUD"), None)
            longitude_column = next((index for index, value in enumerate(names) if value == "LONGITUD"), None)
            entity_reporter_columns = [value for value in names if value.startswith("ENTIDAD REPORTA")]
            identifier_like_columns = [
                value
                for value in names
                if value in {"ID", "CODIGO", "IDENTIFICADOR", "RADICADO", "NUMERO RADICADO"}
                or value.startswith("IDENTIFICADOR ")
                or value.startswith("RADICADO ")
            ]
            record_count = 0
            tolima_count = 0
            missing_date = 0
            missing_municipality = 0
            tolima_missing_area_total = 0
            tolima_missing_latitude = 0
            tolima_missing_longitude = 0
            national_dates: list[str] = []
            tolima_dates: list[str] = []
            tolima_municipalities: set[str] = set()
            tolima_municipality_dates: set[tuple[str, str]] = set()
            for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                department = normalized(row[department_column] if department_column < len(row) else None)
                if not department:
                    continue
                record_count += 1
                date_value = as_iso_date(row[day_column] if day_column < len(row) else None)
                if date_value is None:
                    missing_date += 1
                else:
                    national_dates.append(date_value)
                if department == "TOLIMA":
                    tolima_count += 1
                    municipality = normalized(row[municipality_column] if municipality_column < len(row) else None)
                    if not municipality:
                        missing_municipality += 1
                    else:
                        tolima_municipalities.add(municipality)
                    if area_column is None or row[area_column] is None:
                        tolima_missing_area_total += 1
                    if latitude_column is None or row[latitude_column] is None:
                        tolima_missing_latitude += 1
                    if longitude_column is None or row[longitude_column] is None:
                        tolima_missing_longitude += 1
                    if date_value is not None:
                        tolima_dates.append(date_value)
                        if municipality:
                            tolima_municipality_dates.add((municipality, date_value))
            total_data_rows += record_count
            total_tolima_rows += tolima_count
            sheet_audit = {
                "source_book": source.name,
                "sheet": worksheet.title,
                "worksheet_max_row": worksheet.max_row,
                "worksheet_max_column": worksheet.max_column,
                "header_row": header_row,
                "first_35_headers": [str(value) if value is not None else None for value in header[:35]],
                "record_count_with_department": record_count,
                "national_min_date": min(national_dates) if national_dates else None,
                "national_max_date": max(national_dates) if national_dates else None,
                "missing_date_count": missing_date,
                "tolima_record_count": tolima_count,
                "tolima_min_date": min(tolima_dates) if tolima_dates else None,
                "tolima_max_date": max(tolima_dates) if tolima_dates else None,
                "tolima_unique_date_count": len(set(tolima_dates)),
                "tolima_unique_municipality_count": len(tolima_municipalities),
                "tolima_unique_municipality_date_count": len(tolima_municipality_dates),
                "tolima_missing_municipality_count": missing_municipality,
                "tolima_missing_area_total_count": tolima_missing_area_total,
                "tolima_missing_latitude_count": tolima_missing_latitude,
                "tolima_missing_longitude_count": tolima_missing_longitude,
                "has_area_total_column": bool(area_columns),
                "has_entity_reporter_column": bool(entity_reporter_columns),
                "identifier_like_columns": identifier_like_columns,
            }
            sheets.append(sheet_audit)
            book_sheets.append(sheet_audit)
        books.append(
            {
                "source_file": str(source.relative_to(ROOT)).replace("\\", "/"),
                "sha256": sha256(source),
                "sheet_count": len(book_sheets),
            }
        )
    audit = {
        "source_books": books,
        "sheet_count": len(sheets),
        "total_record_count_with_department": total_data_rows,
        "total_tolima_record_count": total_tolima_rows,
        "sheets": sheets,
        "interpretation_limit": (
            "Los registros son eventos reportados en una base histórica. Esta auditoría no prueba "
            "vigilancia exhaustiva, ausencia de incendios cuando no hay fila, comparabilidad entre años "
            "ni disponibilidad para 2026."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
