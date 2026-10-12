"""Audita cobertura y variables de Tolima en la base agrícola EVA 2019 a 2024."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "eva" / "eva_agricola_2019_2024.xlsx"
OUTPUT = ROOT / "results" / "exploratory" / "inspect_eva_agricola_2019_2024_tolima.json"


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La inspección EVA de Tolima ya existe. No se sobreescriben resultados.")
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["BasePagina"]
    rows = sheet.iter_rows(values_only=True)
    for _ in range(7):
        next(rows)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    department_index = header.index("Departamento")
    year_index = header.index("Año")
    municipality_index = header.index("Municipio")
    crop_index = header.index("Cultivo")
    cycle_index = header.index("Ciclo del cultivo")
    total_records = 0
    tolima_records = 0
    years: Counter[str] = Counter()
    municipalities: Counter[str] = Counter()
    crops: Counter[str] = Counter()
    cycles: Counter[str] = Counter()
    for row in rows:
        if not any(value is not None for value in row):
            continue
        total_records += 1
        if str(row[department_index]).strip().upper() == "TOLIMA":
            tolima_records += 1
            years[str(row[year_index])] += 1
            municipalities[str(row[municipality_index])] += 1
            crops[str(row[crop_index])] += 1
            cycles[str(row[cycle_index])] += 1
    workbook.close()
    audit = {
        "source_file": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
        "total_data_records": total_records,
        "tolima_record_count": tolima_records,
        "tolima_records_by_year": dict(sorted(years.items())),
        "tolima_unique_municipality_count": len(municipalities),
        "tolima_records_by_cycle": dict(sorted(cycles.items())),
        "tolima_top_crops_by_record_count": crops.most_common(20),
        "interpretation_limit": (
            "EVA está agregada por municipio, cultivo y periodo. La ruptura documentada en 2022 para los cultivos "
            "transitorios impide interpretar la serie completa como una medida homogénea sin un diseño específico."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "total_data_records": audit["total_data_records"],
        "tolima_record_count": audit["tolima_record_count"],
        "tolima_records_by_year": audit["tolima_records_by_year"],
        "tolima_unique_municipality_count": audit["tolima_unique_municipality_count"],
        "tolima_records_by_cycle": audit["tolima_records_by_cycle"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
