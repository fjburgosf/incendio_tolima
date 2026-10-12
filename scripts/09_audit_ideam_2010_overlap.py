"""Evalúa el solapamiento potencial de 2010 entre dos libros crudos IDEAM.

No modifica, combina ni deduplica los libros. Una coincidencia exacta de una
firma común es evidencia de posible duplicación, no prueba de identidad de evento.
"""

from __future__ import annotations

import json
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EARLY = ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2002_2010.xlsx"
LATE = ROOT / "data" / "raw" / "ideam" / "bd_icv_nacional_2010_2020.xlsx"
OUTPUT = ROOT / "results" / "exploratory" / "audit_ideam_2010_overlap.json"
SIGNATURE_WIDTH = 35


def normalized(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return " ".join(text.upper().split())


def find_header(worksheet) -> int:
    for number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        values = [normalized(value) for value in row]
        if "DEPARTAMENTO" in values and "MUNICIPIO" in values and any(value.startswith("DIA") for value in values):
            return number
    raise ValueError(f"No se encontró encabezado de datos en {worksheet.title}")


def year_value(value: object) -> int | None:
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, date):
        return value.year
    if isinstance(value, (int, float)) and int(value) == value:
        return int(value)
    return None


def read_2010_records(source: Path, sheet_name: str) -> dict:
    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    header_row = find_header(worksheet)
    header = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    names = [normalized(value) for value in header]
    year_column = names.index("ANO")
    department_column = names.index("DEPARTAMENTO")
    all_signatures: set[tuple[str, ...]] = set()
    tolima_signatures: set[tuple[str, ...]] = set()
    all_rows = 0
    tolima_rows = 0
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        if year_column >= len(row) or year_value(row[year_column]) != 2010:
            continue
        all_rows += 1
        signature = tuple(normalized(value) for value in row[:SIGNATURE_WIDTH])
        all_signatures.add(signature)
        department = normalized(row[department_column] if department_column < len(row) else None)
        if department == "TOLIMA":
            tolima_rows += 1
            tolima_signatures.add(signature)
    return {
        "source_file": str(source.relative_to(ROOT)).replace("\\", "/"),
        "sheet": sheet_name,
        "signature_columns": [str(value) if value is not None else None for value in header[:SIGNATURE_WIDTH]],
        "row_count_2010": all_rows,
        "unique_signature_count_2010": len(all_signatures),
        "tolima_row_count_2010": tolima_rows,
        "tolima_unique_signature_count_2010": len(tolima_signatures),
        "all_signatures": all_signatures,
        "tolima_signatures": tolima_signatures,
    }


def main() -> None:
    early = read_2010_records(EARLY, "BD_ICV_2002_2018_V_0_2019")
    late = read_2010_records(LATE, "BD_ICV_2010_2020")
    common_all = early["all_signatures"] & late["all_signatures"]
    common_tolima = early["tolima_signatures"] & late["tolima_signatures"]
    audit = {
        "comparison_year": 2010,
        "signature_definition": (
            "Valores normalizados de las primeras 35 columnas de datos, en el orden del encabezado. "
            "No es un identificador oficial de evento."
        ),
        "early_book": {key: value for key, value in early.items() if not key.endswith("signatures")},
        "late_book": {key: value for key, value in late.items() if not key.endswith("signatures")},
        "exact_common_signature_count_2010": len(common_all),
        "exact_common_tolima_signature_count_2010": len(common_tolima),
        "interpretation_limit": (
            "El solapamiento temporal exige no combinar los libros sin una regla documentada. "
            "La ausencia de una firma común no descarta registros del mismo evento con campos distintos."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
