"""Audita el enlace nominal de las bases EVA con los municipios de Cortolima."""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOUNDARIES = ROOT / "data" / "raw" / "boundaries" / "cortolima_limite_municipal_25k.geojson"
SOURCES = {
    "agricola": (ROOT / "data" / "raw" / "eva" / "eva_agricola_2019_2024.xlsx", "BasePagina", 8),
    "pecuaria": (ROOT / "data" / "raw" / "eva" / "eva_pecuaria_2019_2024.xlsx", None, None),
}
OUTPUT = ROOT / "results" / "exploratory" / "audit_eva_cortolima_municipal_linkage.json"


def normalized(value: object) -> str:
    value = "" if value is None else str(value)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]+", "", value.upper())


def header_at(sheet, row: int) -> list[str]:
    return [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=row, max_row=row, values_only=True))]


def locate_header(sheet) -> tuple[int, list[str]]:
    for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        header = [str(value).strip() if value is not None else "" for value in row]
        if {"Departamento", "Municipio", "Código Dane municipio"}.issubset(header):
            return number, header
    raise ValueError(f"No se encontró encabezado EVA en {sheet.title}")


def audit_sheet(sheet, header_row: int, boundary_names: dict[str, str]) -> dict[str, object]:
    header = header_at(sheet, header_row)
    dept_index = header.index("Departamento")
    municipality_index = header.index("Municipio")
    code_index = header.index("Código Dane municipio")
    names: dict[str, set[str]] = {}
    codes: set[str] = set()
    record_count = 0
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in row):
            continue
        if normalized(row[dept_index]) != "TOLIMA":
            continue
        record_count += 1
        municipality = str(row[municipality_index]).strip()
        names.setdefault(normalized(municipality), set()).add(municipality)
        code = str(row[code_index]).strip()
        if code:
            codes.add(code.zfill(5))
    unresolved = sorted(
        {original for key, originals in names.items() if key not in boundary_names for original in originals},
        key=normalized,
    )
    return {
        "sheet": sheet.title,
        "tolima_record_count": record_count,
        "unique_tolima_municipality_name_count": len(names),
        "unique_tolima_municipality_code_count": len(codes),
        "unique_direct_normalized_match_count": sum(key in boundary_names for key in names),
        "unresolved_eva_municipalities": unresolved,
    }


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La auditoría de enlace EVA ya existe. No se sobreescriben resultados.")
    collection = json.loads(BOUNDARIES.read_text(encoding="utf-8"))
    boundary_names = {
        normalized(feature["properties"].get("mun_nombre")): feature["properties"].get("mun_nombre")
        for feature in collection["features"]
    }
    summaries = []
    agricultural, sheet_name, header_row = SOURCES["agricola"]
    workbook = load_workbook(agricultural, read_only=True, data_only=True)
    summaries.append({"source": agricultural.name, **audit_sheet(workbook[sheet_name], header_row, boundary_names)})
    workbook.close()
    livestock, _, _ = SOURCES["pecuaria"]
    workbook = load_workbook(livestock, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        found_row, _ = locate_header(sheet)
        summaries.append({"source": livestock.name, **audit_sheet(sheet, found_row, boundary_names)})
    workbook.close()
    audit = {
        "boundary_source": str(BOUNDARIES.relative_to(ROOT)).replace("\\", "/"),
        "boundary_unique_municipality_name_count": len(boundary_names),
        "name_normalization": "Mayúsculas sin tildes ni caracteres no alfanuméricos",
        "sheets": summaries,
        "interpretation_limit": (
            "La coincidencia por nombre valida solo una llave administrativa potencial. No localiza predios, animales, "
            "cultivos, daños o eventos de incendio individuales. Los códigos DANE requieren validación contra una fuente DIVIPOLA antes de ser una llave primaria."
        ),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
