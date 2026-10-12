"""Descarga y audita la base agrícola EVA 2019 a 2024 publicada por UPRA."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from urllib.request import urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
URL = "https://upra.gov.co/sites/default/files/2025-07/20250617_BaseAgricola20192024.xlsx"
RAW_DIR = ROOT / "data" / "raw" / "eva"
OUTPUT = RAW_DIR / "eva_agricola_2019_2024.xlsx"
MANIFEST = RAW_DIR / "eva_agricola_2019_2024_manifest.json"
AUDIT = ROOT / "results" / "exploratory" / "audit_eva_agricola_2019_2024.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def first_nonempty_rows(sheet, limit: int = 8) -> list[list[object]]:
    result: list[list[object]] = []
    for row in sheet.iter_rows(values_only=True):
        values = list(row)
        if any(value is not None for value in values):
            result.append(values)
        if len(result) >= limit:
            break
    return result


def main() -> None:
    if OUTPUT.exists() or MANIFEST.exists() or AUDIT.exists():
        raise FileExistsError("La auditoría EVA ya existe. No se sobreescriben insumos ni resultados.")
    with urlopen(URL, timeout=180) as response:
        payload = response.read()
        content_type = response.headers.get_content_type()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_bytes(payload)
    manifest = {
        "url": URL,
        "file": OUTPUT.name,
        "bytes": len(payload),
        "sha256": sha256(OUTPUT),
        "content_type": content_type,
        "purpose": "Auditoría de estructura EVA. No constituye un análisis de impacto o causalidad.",
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    workbook = load_workbook(OUTPUT, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "first_nonempty_rows": first_nonempty_rows(sheet),
        })
    workbook.close()
    audit = {
        "source_file": str(OUTPUT.relative_to(ROOT)).replace("\\", "/"),
        "sheet_count": len(sheets),
        "sheets": sheets,
        "interpretation_limit": (
            "La fuente está agregada por municipio, cultivo y periodo. No ubica predios ni eventos de incendio "
            "y no permite atribuir cambios productivos a un episodio."
        ),
    }
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps(audit, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({
        "bytes": manifest["bytes"],
        "sha256": manifest["sha256"],
        "sheet_count": audit["sheet_count"],
        "sheets": [{"name": row["name"], "max_row": row["max_row"], "max_column": row["max_column"]} for row in sheets],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
