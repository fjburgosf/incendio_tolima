"""Audita el enlace nominal MapBiomas con municipios de Cortolima."""

from __future__ import annotations

import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "mapbiomas" / "mapbiomas_colombia_collection3_coverage_statistics.zip"
BOUNDARIES = ROOT / "data" / "raw" / "boundaries" / "cortolima_limite_municipal_25k.geojson"
CROSSWALK = ROOT / "configs" / "ideam_cortolima_municipality_crosswalk.json"
OUTPUT = ROOT / "results" / "exploratory" / "audit_mapbiomas_cortolima_municipal_linkage.json"


def normalized(value: object) -> str:
    value = "" if value is None else str(value)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]+", "", value.upper())


def main() -> None:
    if OUTPUT.exists():
        raise FileExistsError("La auditoría de enlace MapBiomas ya existe. No se sobreescriben resultados.")
    boundaries = json.loads(BOUNDARIES.read_text(encoding="utf-8"))
    boundary_names = {normalized(feature["properties"].get("mun_nombre")) for feature in boundaries["features"]}
    crosswalk = json.loads(CROSSWALK.read_text(encoding="utf-8"))["mappings"]
    with zipfile.ZipFile(SOURCE) as archive, archive.open("Statistics-for-Website-MB-Cobertura-col3.xlsx") as binary:
        workbook = load_workbook(io.BytesIO(binary.read()), read_only=True, data_only=True)
        sheet = workbook["COBERTURA_MUNICIPIO"]
        rows = sheet.iter_rows(values_only=True)
        header = [str(value).strip() if value is not None else "" for value in next(rows)]
        dept_index = header.index("departamento")
        municipality_index = header.index("municipio")
        names = {str(row[municipality_index]).strip() for row in rows if str(row[dept_index]).strip().upper() == "TOLIMA"}
        workbook.close()
    direct = sorted(name for name in names if normalized(name) in boundary_names)
    crosswalk_matches = sorted(name for name in names if normalized(name) not in boundary_names and normalized(name) in crosswalk)
    unresolved = sorted(name for name in names if normalized(name) not in boundary_names and normalized(name) not in crosswalk)
    audit = {
        "source_file": str(SOURCE.relative_to(ROOT)).replace("\\", "/"),
        "boundary_source": str(BOUNDARIES.relative_to(ROOT)).replace("\\", "/"),
        "crosswalk_source": str(CROSSWALK.relative_to(ROOT)).replace("\\", "/"),
        "unique_mapbiomas_tolima_municipality_count": len(names),
        "direct_normalized_match_count": len(direct),
        "crosswalk_match_count": len(crosswalk_matches),
        "crosswalk_mapbiomas_municipalities": crosswalk_matches,
        "unresolved_mapbiomas_municipalities": unresolved,
        "interpretation_limit": "La coincidencia nominal solo permite un enlace administrativo potencial. No identifica píxeles, predios ni procesos de cambio de cobertura individuales.",
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
